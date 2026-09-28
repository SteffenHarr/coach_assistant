"""Export builders — Excel/PDF timetable for a single plan, Excel for the
player and coach rosters. Pure functions (bytes in, bytes out); routers.py
wires these to actual HTTP endpoints (planner/admin only, see there)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from coach_api.domain.time_grid import TimeGrid

WEEKDAYS_DE = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]

_BRAND = "2563EB"
_DAY_BAND = "1E3A5F"
_STRIPE = "F2F4F7"
_COACH_BAND = "DBEAFE"


def _slot_label(local_idx: int, grid: TimeGrid) -> str:
    minutes = grid.day_start_minutes + local_idx * grid.slot_minutes
    h, m = divmod(minutes, 60)
    return f"{h:02d}:{m:02d}"


def _used_court_ids(sessions: list[dict], court_ids: list[str]) -> list[str]:
    """Courts that never appear in any session are dropped entirely — no
    point printing/exporting a column that's always empty."""
    used = {str(s.get("court_id")) for s in sessions}
    return [cid for cid in court_ids if cid in used]


def _day_blocks(
    sessions: list[dict],
    day: int,
    court_ids: list[str],
    coach_names: dict[str, str],
    player_names: dict[str, str],
    grid: TimeGrid,
) -> dict[str, dict[int, tuple[int, str, str]]]:
    """court_id -> {start_local_slot: (duration_slots, coach, content)} for
    one weekday. Coach and content (the free-text label if set, e.g. a team
    like "Damen 30", otherwise the comma-joined player names) are kept apart
    so the builders can render the coach clearly distinct from the
    players/students (bold, colored) instead of just stacking plain text."""
    spd = grid.slots_per_day
    by_court: dict[str, dict[int, tuple[int, str, str]]] = {cid: {} for cid in court_ids}
    for s in sessions:
        idxs = sorted(s.get("slot_indices") or [])
        if not idxs or idxs[0] // spd != day:
            continue
        cid = str(s.get("court_id"))
        if cid not in by_court:
            by_court[cid] = {}
        start_local = idxs[0] % spd
        # Haupttrainer plus manuell ergänzte Zusatztrainer.
        coach = " + ".join(
            coach_names.get(str(coach_id), "?")
            for coach_id in [s.get("coach_id"), *(s.get("extra_coach_ids") or [])]
        )
        note = s.get("note")
        if note:
            coach = f"{coach} · {note}"
        label = s.get("label")
        pnotes = s.get("player_notes") or {}

        def _named(pid) -> str:
            name = player_names.get(str(pid), "?")
            extra = pnotes.get(str(pid))
            return f"{name} ({extra})" if extra else name

        content = label if label else ", ".join(_named(p) for p in (s.get("player_ids") or []))
        by_court[cid][start_local] = (len(idxs), coach, content)
    return by_court


def _used_time_range(by_court: dict[str, dict[int, tuple[int, str, str]]]) -> tuple[int, int] | None:
    """(min_local_slot, max_local_slot) actually covered by any block that
    day, inclusive — so leading/trailing hours nobody ever uses aren't
    printed. Returns None if the day has no sessions at all."""
    lo = None
    hi = None
    for blocks in by_court.values():
        for start, (dur, _coach, _content) in blocks.items():
            end = start + dur - 1
            lo = start if lo is None else min(lo, start)
            hi = end if hi is None else max(hi, end)
    if lo is None or hi is None:
        return None
    return lo, hi


def build_plan_xlsx(
    title: str,
    sessions: list[dict],
    coach_names: dict[str, str],
    court_ids: list[str],
    court_names: dict[str, str],
    player_names: dict[str, str],
    grid: TimeGrid | None = None,
) -> bytes:
    """Ein einziges Blatt statt sieben Reitern: alle Wochentage kompakt
    untereinander gestapelt, jeweils mit eigener Überschriftenzeile, damit
    man nicht für jeden Tag einen Tab wechseln muss. Layout pro Tag wie
    gehabt: Zeit runter, Plätze quer (wie das Kalenderraster in der App).

    Jeder Termin bekommt eine eigene schmale Trainer-Kopfzeile (fett,
    farbig) und darunter eine eigene, größere Zelle für Spieler/Freitext —
    zwei physisch getrennte Zellen statt einer gemeinsamen Zelle mit
    In-Cell-Rich-Text. Rich-Text-Zellen führten bei manchen Excel-Versionen
    beim Öffnen zu "Datei muss repariert werden" und gingen dabei ihre
    Formatierung verlieren — mit echten, separaten Zellen passiert das
    nicht."""
    grid = grid or TimeGrid()
    court_ids = _used_court_ids(sessions, court_ids)
    ncols = 1 + len(court_ids)

    wb = Workbook()
    ws = wb.active
    ws.title = "Trainingsplan"

    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([])

    any_day = False
    for day in range(7):
        by_court = _day_blocks(sessions, day, court_ids, coach_names, player_names, grid)
        time_range = _used_time_range(by_court)
        if time_range is None:
            continue
        any_day = True
        lo, hi = time_range
        n_rows = hi - lo + 1

        ws.append([WEEKDAYS_DE[day]])
        day_row = ws.max_row
        ws.merge_cells(start_row=day_row, start_column=1, end_row=day_row, end_column=ncols)
        day_cell = ws.cell(row=day_row, column=1)
        day_cell.font = Font(bold=True, color="FFFFFF", size=12)
        day_cell.fill = PatternFill("solid", fgColor=_DAY_BAND)
        day_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[day_row].height = 22

        ws.append(["Zeit"] + [court_names.get(cid, "?") for cid in court_ids])
        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=_BRAND)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Zeitspalte für alle Zeilen dieses Tages anlegen. Bewusst über
        # direkte Zell-Zugriffe statt ws.append(), damit die Zeilennummern
        # unten beim Einzeichnen der Termine (die über den Blockstart
        # hinausreichen) feststehen.
        first_data_row = header_row + 1
        for i in range(n_rows):
            r = first_data_row + i
            time_cell = ws.cell(row=r, column=1, value=_slot_label(lo + i, grid))
            time_cell.alignment = Alignment(horizontal="right", vertical="center")
            ws.row_dimensions[r].height = 26

        for col_i, cid in enumerate(court_ids, start=2):
            for start, (dur, coach, content) in by_court.get(cid, {}).items():
                coach_row = first_data_row + (start - lo)
                coach_cell = ws.cell(row=coach_row, column=col_i, value=coach)
                coach_cell.font = Font(bold=True, color=_BRAND, size=9)
                coach_cell.alignment = Alignment(horizontal="center", vertical="center")
                coach_cell.fill = PatternFill("solid", fgColor=_COACH_BAND)

                if dur < 2:
                    # Sollte durch die Mindestdauer des Solvers (60 Min.)
                    # nicht vorkommen — zur Sicherheit trotzdem abgedeckt.
                    coach_cell.value = f"{coach} – {content}"
                    continue

                content_row0 = coach_row + 1
                content_row1 = coach_row + dur - 1
                content_cell = ws.cell(row=content_row0, column=col_i, value=content)
                content_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                content_cell.fill = PatternFill("solid", fgColor=_STRIPE)
                if content_row1 > content_row0:
                    ws.merge_cells(start_row=content_row0, start_column=col_i, end_row=content_row1, end_column=col_i)

        ws.append([])  # kompakter Trenner zum nächsten Tag

    if not any_day:
        ws.append([f"{title} — keine Einheiten in diesem Plan."])

    ws.column_dimensions["A"].width = 8
    for i in range(len(court_ids)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 24
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_plan_pdf(
    title: str,
    sessions: list[dict],
    coach_names: dict[str, str],
    court_ids: list[str],
    court_names: dict[str, str],
    player_names: dict[str, str],
    grid: TimeGrid | None = None,
) -> bytes:
    """Ein durchlaufendes Dokument statt einer festen Seite pro Wochentag —
    kurze Tage teilen sich eine Seite, nur wenn's nicht mehr passt, gibt's
    einen automatischen Seitenumbruch. Gleiches Rasterlayout wie der
    Excel-Export und der In-App-Kalender."""
    grid = grid or TimeGrid()
    court_ids = _used_court_ids(sessions, court_ids)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), title=title,
        topMargin=10 * mm, bottomMargin=10 * mm, leftMargin=8 * mm, rightMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "cell", parent=styles["Normal"], fontName="Helvetica", fontSize=6.5, leading=7.5, alignment=TA_CENTER,
    )
    # Trainer klar von Spielern/Schülern getrennt: fett + Markenfarbe oben,
    # Spieler/Freitext normal darunter.
    coach_span = f"<b><font color='#{_BRAND}'>%s</font></b>"

    elements: list = [Paragraph(title, styles["Title"]), Spacer(1, 4 * mm)]
    any_day = False

    for day in range(7):
        by_court = _day_blocks(sessions, day, court_ids, coach_names, player_names, grid)
        time_range = _used_time_range(by_court)
        if time_range is None:
            continue
        lo, hi = time_range
        any_day = True

        day_elements: list = [Paragraph(WEEKDAYS_DE[day], styles["Heading2"]), Spacer(1, 2 * mm)]

        header = ["Zeit"] + [court_names.get(cid, "?") for cid in court_ids]
        data = [header]
        row_of_slot: dict[int, int] = {}
        for local in range(lo, hi + 1):
            row_of_slot[local] = len(data)
            row = [_slot_label(local, grid)]
            for cid in court_ids:
                block = by_court.get(cid, {}).get(local)
                if block:
                    _dur, coach, content = block
                    row.append(Paragraph(f"{coach_span % coach}<br/>{content}", cell_style))
                else:
                    row.append("")
            data.append(row)

        span_cmds: list[tuple] = []
        bg_cmds: list[tuple] = []
        for col_i, cid in enumerate(court_ids, start=1):
            for start, (dur, _coach, _content) in by_court.get(cid, {}).items():
                r0 = row_of_slot[start]
                bg_cmds.append(("BACKGROUND", (col_i, r0), (col_i, r0), colors.HexColor(f"#{_STRIPE}")))
                if dur > 1:
                    r1 = row_of_slot[start + dur - 1]
                    span_cmds.append(("SPAN", (col_i, r0), (col_i, r1)))

        table = Table(data, repeatRows=1, colWidths=[16 * mm] + [None] * len(court_ids))
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_BRAND}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 6.5),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ] + span_cmds + bg_cmds
        table.setStyle(TableStyle(style))
        day_elements.append(table)
        day_elements.append(Spacer(1, 5 * mm))

        # Hält Tagesüberschrift + Tabelle zusammen, erzwingt aber keinen
        # Seitenumbruch — mehrere kurze Tage landen so auf einer Seite.
        elements.append(KeepTogether(day_elements))

    if not any_day:
        elements = [Paragraph(f"{title} — keine Einheiten in diesem Plan.", styles["Title"])]

    doc.build(elements)
    return buf.getvalue()


def _compute_age(birth_date_iso: str | None) -> int | None:
    if not birth_date_iso:
        return None
    try:
        b = date.fromisoformat(birth_date_iso)
    except ValueError:
        return None
    today = date.today()
    return today.year - b.year - ((today.month, today.day) < (b.month, b.day))


def _player_rows(players: list[dict]) -> list[list[str]]:
    out = []
    for p in players:
        pref = p.get("preferences") or {}
        age = _compute_age(pref.get("birth_date")) if pref.get("birth_date") else pref.get("age")
        out.append([
            p.get("name", ""),
            pref.get("birth_date") or "",
            str(age) if age is not None else "",
            str(pref.get("level_lk")) if pref.get("level_lk") is not None else "",
            ", ".join(p.get("categories") or []),
            str(p.get("min_slots_per_week", "")),
            str(p.get("max_slots_per_week", "")),
            "Ja" if p.get("has_account") else "Nein",
            "Ja" if p.get("availability") else "Nein",
            pref.get("notes") or "",
        ])
    return out


def _coach_rows(coaches: list[dict]) -> list[list[str]]:
    out = []
    for c in coaches:
        out.append([
            c.get("name", ""),
            ", ".join(c.get("categories") or []),
            str(c.get("max_group_size", "")),
            "Ja" if c.get("has_account") else "Nein",
        ])
    return out


PLAYER_HEADER = [
    "Name", "Geburtsdatum", "Alter", "LK", "Kategorien", "Min Std/Woche", "Max Std/Woche", "Konto",
    "Verfügbarkeit eingetragen", "Kommentar",
]
COACH_HEADER = ["Name", "Kategorien", "Max Gruppengröße", "Konto"]


def _write_sheet(ws, header: list[str], rows: list[list[str]]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_BRAND)
    for i, row in enumerate(rows):
        ws.append(row)
        if i % 2 == 1:
            for cell in ws[2 + i]:
                cell.fill = PatternFill("solid", fgColor=_STRIPE)
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 45)
    ws.freeze_panes = "A2"


def build_players_xlsx(players: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spieler"
    _write_sheet(ws, PLAYER_HEADER, _player_rows(players))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_coaches_xlsx(coaches: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Trainer"
    _write_sheet(ws, COACH_HEADER, _coach_rows(coaches))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_people_xlsx(players: list[dict], coaches: list[dict]) -> bytes:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Spieler"
    _write_sheet(ws1, PLAYER_HEADER, _player_rows(players))
    ws2 = wb.create_sheet("Trainer")
    _write_sheet(ws2, COACH_HEADER, _coach_rows(coaches))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()




# Mindestanzahl Termin-Spaltenpaare je Spieler. Hat jemand mehr Einheiten,
# wächst die Tabelle automatisch mit, damit nie ein Termin unter den Tisch
# fällt — im Normalfall bleibt es bei drei.
_MIN_SCHEDULE_SLOTS = 3


def build_player_schedule_xlsx(
    title: str,
    sessions: list[dict],
    coach_names: dict[str, str],
    court_names: dict[str, str],
    player_names: dict[str, str],
    grid: TimeGrid | None = None,
) -> bytes:
    """Spieler-Sicht auf den Plan: eine Zeile je Spieler (alphabetisch),
    daneben die Trainingseinheiten nebeneinander — je Einheit ein
    Spaltenpaar aus Beschreibung und Termin (Wochentag, Uhrzeit, Trainer).

    Gedacht zum Weitergeben an Eltern/Spieler — im Gegensatz zum
    Stundenplan-Export (Zeit runter, Plätze quer) findet man hier den
    eigenen Namen sofort und liest in einer Zeile ab, wann man dran ist.
    Der Platz steht bewusst nicht dabei: er interessiert beim Nachschlagen
    der eigenen Zeiten nicht und kostet nur Breite.

    Freitext-Einheiten (z.B. "Damen 30") tauchen nicht auf, da sie keine
    einzelnen Spieler hinterlegt haben.
    """
    grid = grid or TimeGrid()
    spd = grid.slots_per_day
    short_day = [d[:2] for d in WEEKDAYS_DE]

    def _time_range(idxs: list[int]) -> str:
        start = idxs[0] % spd
        end_minutes = grid.day_start_minutes + (idxs[-1] % spd + 1) * grid.slot_minutes
        h, m = divmod(end_minutes, 60)
        return f"{_slot_label(start, grid)}–{h:02d}:{m:02d}"

    # player_id -> Liste von (sortier_schluessel, beschreibung, termin)
    per_player: dict[str, list[tuple[int, str, str]]] = {}
    for s in sessions:
        idxs = sorted(s.get("slot_indices") or [])
        pids = s.get("player_ids") or []
        if not idxs or not pids:
            continue
        day = idxs[0] // spd
        coach = " + ".join(
            coach_names.get(str(cid), "?")
            for cid in [s.get("coach_id"), *(s.get("extra_coach_ids") or [])]
        )
        termin = f"{short_day[day]} {_time_range(idxs)}, {coach}"
        note = s.get("note")
        pnotes = s.get("player_notes") or {}
        for pid in pids:
            # Beschreibung der Einheit (z.B. "U15") und die persönliche
            # Notiz des Spielers (z.B. "gerade Wochen") zusammen in einer
            # Zelle — beides beschreibt denselben Termin.
            parts = [d for d in (note, pnotes.get(str(pid))) if d]
            per_player.setdefault(str(pid), []).append((idxs[0], ", ".join(parts), termin))

    for entries in per_player.values():
        entries.sort(key=lambda e: e[0])   # chronologisch über die Woche

    n_units = max(_MIN_SCHEDULE_SLOTS, max((len(v) for v in per_player.values()), default=0))
    ncols = 2 + 2 * n_units

    wb = Workbook()
    ws = wb.active
    ws.title = "Spieler-Zeiten"

    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([])

    header = ["Spieler", "Std/Woche"]
    for i in range(1, n_units + 1):
        header += [f"Beschreibung {i}", f"Training {i}"]
    ws.append(header)
    header_row = ws.max_row

    ordered = sorted(per_player.items(), key=lambda kv: player_names.get(kv[0], "").lower())
    rows: list[list] = []
    for pid, entries in ordered:
        slots_total = sum(
            len(s.get("slot_indices") or [])
            for s in sessions
            if str(pid) in [str(p) for p in (s.get("player_ids") or [])]
        )
        hours = slots_total * grid.slot_minutes / 60
        row_values: list = [player_names.get(pid, "?"), round(hours, 1)]
        for j in range(n_units):
            if j < len(entries):
                row_values += [entries[j][1], entries[j][2]]
            else:
                row_values += ["", ""]
        rows.append(row_values)
        ws.append(row_values)

    if not ordered:
        ws.append([f"{title} — keine Spieler-Einheiten in diesem Plan."])

    # Spaltenbreiten aus dem tatsächlichen Inhalt: sonst wird beim Drucken
    # bzw. beim PDF-Export abgeschnitten, was die Nachbarzelle verdeckt.
    # Zusätzlich steht überall Zeilenumbruch an und die Zeilenhöhe bleibt
    # ungesetzt, damit Excel/LibreOffice automatisch höher machen —
    # dadurch kann auch bei ungewöhnlich langen Namen nichts wegfallen.
    _CAP = 34
    for col in range(1, ncols + 1):
        longest = max(
            [len(str(header[col - 1]))]
            + [len(str(r[col - 1])) for r in rows if r[col - 1] not in (None, "")]
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 3, 9), _CAP)

    # Klare Trennung: dünnes Gitter überall, kräftige Linie zwischen den
    # Blöcken (Name/Stunden und je Trainingseinheit), damit man beim
    # Querlesen die Einheiten nicht verwechselt.
    thin = Side(style="thin", color="C7CDD6")
    strong = Side(style="medium", color="8A94A6")
    block_starts = {3 + 2 * j for j in range(n_units)}   # jede Beschreibung-Spalte
    for r in range(header_row, ws.max_row + 1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = Border(
                left=strong if col in block_starts or col == 2 else thin,
                right=strong if col == ncols else thin,
                top=thin,
                bottom=thin,
            )
            if r == header_row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=_BRAND)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                continue
            cell.alignment = Alignment(
                horizontal="center" if col == 2 else "left",
                vertical="center",
                wrap_text=True,
            )
            if col == 1:
                cell.font = Font(bold=True)
            if (r - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=_STRIPE)

    # Kopfzeile sowie Name und Stunden beim Scrollen stehen lassen.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    # Druckbild: die Liste wird ausgedruckt und ausgehängt, deshalb müssen
    # *alle* Spalten auf eine Seitenbreite passen — quer, auf Breite
    # skaliert, aber ohne Höhenbegrenzung: nach unten darf sie über
    # beliebig viele Seiten laufen.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for m in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, m, 0.4)
    ws.print_area = f"A1:{get_column_letter(ncols)}{ws.max_row}"
    # Überschriftenzeile auf jeder Folgeseite wiederholen.
    ws.print_title_rows = f"{header_row}:{header_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
